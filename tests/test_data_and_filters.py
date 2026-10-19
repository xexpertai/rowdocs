"""Data loading, header normalisation, row selection, filters and file-name safety."""

from __future__ import annotations

from datetime import datetime

import pytest
from conftest import write_csv
from openpyxl import Workbook

from rowdocs.data import DataError, load_csv_bytes, load_data, normalise_header, parse_rows, select_rows
from rowdocs.filters import date_filter, money
from rowdocs.naming import MAX_NAME_LENGTH, plan_names, sanitize


def test_header_normalisation_and_mapping():
    assert normalise_header("  Full Name ") == "full_name"
    assert normalise_header("Unit Price (₹)") == "unit_price"
    assert normalise_header("E-mail / Phone") == "e_mail_phone"
    assert normalise_header("2026 Total") == "col_2026_total"
    data = load_csv_bytes(b"Full Name,full name,,Amount\nAsha,A,,10\n")
    assert [(c.header, c.key) for c in data.columns] == [
        ("Full Name", "full_name"),
        ("full name", "full_name_2"),
        ("Amount", "amount"),
    ]


@pytest.mark.parametrize("delimiter", [",", ";", "\t"])
def test_csv_delimiter_detection_and_bom(delimiter):
    raw = ("﻿" + delimiter.join(["Name", "City"]) + "\nAsha" + delimiter + "Pune\n").encode("utf-8")
    data = load_csv_bytes(raw)
    assert data.keys == ["name", "city"]
    assert data.rows == [{"name": "Asha", "city": "Pune"}]


def test_csv_rejects_non_utf8():
    with pytest.raises(DataError, match="UTF-8"):
        load_csv_bytes("Name\nJosé\n".encode("cp1252"))


def test_xlsx_named_sheet_and_blank_rows(tmp_path):
    workbook = Workbook()
    workbook.active.append(["ignored"])
    sheet = workbook.create_sheet("People")
    sheet.append(["Name", "Joined"])
    sheet.append(["Asha", datetime(2026, 3, 5)])
    sheet.append([None, None])
    sheet.append(["Ben", datetime(2026, 4, 1)])
    workbook.save(tmp_path / "p.xlsx")
    data = load_data(tmp_path / "p.xlsx", sheet="People")
    assert [r["name"] for r in data.rows] == ["Asha", "Ben"]
    with pytest.raises(DataError, match="not found"):
        load_data(tmp_path / "p.xlsx", sheet="Nope")


def test_row_selection_parsing():
    assert parse_rows("2-4,7, 9", 10) == [2, 3, 4, 7, 9]
    for bad in ("0", "4-2", "x", "11"):
        with pytest.raises(DataError):
            parse_rows(bad, 10)


def test_where_filters(tmp_path):
    path = write_csv(tmp_path / "d.csv", ["Name", "Status"], [["A", "Paid"], ["B", "unpaid"], ["C", "paid"]])
    assert select_rows(load_data(path), where=["status=paid"]).selected == [1, 3]
    assert select_rows(load_data(path), where=["Status != paid"]).selected == [2]
    assert select_rows(load_data(path), rows="2-3", where=["status=paid"]).selected == [3]
    with pytest.raises(DataError, match="unknown column"):
        select_rows(load_data(path), where=["colour=red"])


def test_money_filter():
    assert money(1234567.5, "₹") == "₹1,234,567.50"
    assert money("1234567.5", "₹") == "₹1,234,567.50"
    assert money("₹1,200", "₹") == "₹1,200.00"
    assert money(-42, "$", 0) == "-$42"
    assert money(2.675, "", 2) == "2.68"  # Decimal rounding, not binary float surprises
    assert money("", "₹") == ""
    with pytest.raises(ValueError, match="money"):
        money("twelve", "₹")


def test_date_filter_excel_cell_iso_string_and_serial(tmp_path):
    workbook = Workbook()
    workbook.active.append(["Issued"])
    workbook.active.append([datetime(2026, 3, 5)])
    workbook.save(tmp_path / "d.xlsx")
    excel_value = load_data(tmp_path / "d.xlsx").rows[0]["issued"]
    assert isinstance(excel_value, datetime)
    assert date_filter(excel_value, "%d %b %Y") == "05 Mar 2026"
    assert date_filter("2026-03-05", "%d %b %Y") == "05 Mar 2026"
    assert date_filter(46086, "%d %b %Y") == "05 Mar 2026"  # Excel serial number
    assert date_filter("", "%d %b %Y") == ""
    with pytest.raises(ValueError, match="date"):
        date_filter("next Tuesday")


def test_sanitize_blocks_traversal_and_limits_length():
    assert sanitize("../../etc/passwd", "row_1") == "_.._etc_passwd"
    assert sanitize("..", "row_1") == "row_1"
    assert sanitize('a<b>:c"|?*', "x") == "a_b__c____"
    assert sanitize("CON", "x") == "_CON"
    assert len(sanitize("x" * 500, "row_1")) == MAX_NAME_LENGTH


def test_plan_names_suffixes_and_reserved_names():
    planned = plan_names("{{ n }}", [(1, {"n": "X"}), (2, {"n": "x"}), (3, {"n": "X"}), (4, {"n": "ALL"})])
    assert [p.final for p in planned] == ["X", "x_2", "X_3", "ALL_2"]
