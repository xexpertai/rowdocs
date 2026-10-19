"""End-to-end CLI tests mapped to the acceptance criteria in SPEC.md section 4."""

from __future__ import annotations

import csv
import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from conftest import docx_xml, make_template, pdf_works, write_csv
from docx import Document
from pypdf import PdfReader

from rowdocs.cli import main

CERT_NAMES = [
    "Asha Verma",
    "Tomas Lindqvist",
    "Mei-Ling Ong",
    "Kwame Mensah",
    "Lucia Ferraro",
    "Ravi Kulkarni",
    "Hana Kobayashi",
    "Daniel O'Connor",
    "Fatima Zahra Idrissi",
    "Pedro Alves",
    "Nadia Petrova",
    "Samuel Adeyemi",
]


def manifest(outdir: Path) -> list[dict[str, str]]:
    with open(outdir / "manifest.csv", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


# Criterion 1 -------------------------------------------------------------------------------------------------
def test_help_lists_subcommands():
    exe = shutil.which("rowdocs", path=str(Path(sys.executable).parent))
    command = [exe, "--help"] if exe else [sys.executable, "-m", "rowdocs.cli", "--help"]
    out = subprocess.run(command, capture_output=True, text=True, check=True).stdout
    for sub in ("check", "build", "serve"):
        assert sub in out


# Criterion 2 -------------------------------------------------------------------------------------------------
def test_check_sample_invoices_warns_about_missing_amount(examples, capsys):
    code = main(["check", str(examples / "invoices.csv"), str(examples / "invoice.docx")])
    out = capsys.readouterr().out
    assert code == 0
    warn_lines = [line for line in out.splitlines() if line.startswith("WARN")]
    assert any("row 5" in line and "amount" in line for line in warn_lines), out
    assert "Invoice No" in out and "invoice_no" in out  # header -> key mapping is shown


def test_check_json_and_strict(examples, capsys):
    args = ["check", str(examples / "invoices.csv"), str(examples / "invoice.docx")]
    assert main([*args, "--json"]) == 0
    report = json.loads(capsys.readouterr().out)
    assert report["ok"] is True and report["errors"] == 0
    empties = [(i["level"], i["row"], i["variable"]) for i in report["issues"] if i["code"] == "empty_value"]
    assert empties == [("WARN", 5, "amount")]
    assert "header" in report["columns"][0] and "footer" in report["variables"]["data_note"]
    assert main([*args, "--strict"]) == 2


# Criterion 3 -------------------------------------------------------------------------------------------------
def test_unknown_field_is_an_error_and_blocks_build(tmp_path, capsys):
    data = write_csv(tmp_path / "d.csv", ["Name"], [["Asha"]])
    template = make_template(tmp_path / "t.docx", ["Hello {{ name }}, {{ unknown_field }}"])
    assert main(["check", str(data), str(template)]) == 2
    out = capsys.readouterr().out
    assert any(line.startswith("ERROR") and "unknown_field" in line for line in out.splitlines())

    outdir = tmp_path / "out"
    assert main(["build", str(data), str(template), "-o", str(outdir)]) == 2
    assert "unknown_field" in capsys.readouterr().err
    assert not outdir.exists()

    assert main(["build", str(data), str(template), "-o", str(outdir), "--force"]) == 0
    assert (outdir / "row_1.docx").is_file()


def test_template_syntax_error_is_reported_not_crashed(tmp_path, capsys):
    data = write_csv(tmp_path / "d.csv", ["Name"], [["Asha"]])
    template = make_template(tmp_path / "t.docx", ["Hello {{ name }", "{{ name|nosuchfilter }}"])
    assert main(["check", str(data), str(template)]) == 2
    assert "template" in capsys.readouterr().out.lower()


# Criterion 4 -------------------------------------------------------------------------------------------------
def test_build_certificates(examples, tmp_path):
    outdir = tmp_path / "out"
    code = main(
        [
            "build",
            str(examples / "certificates.xlsx"),
            str(examples / "certificate.docx"),
            "-o",
            str(outdir),
            "--name",
            "{{ full_name }}",
        ]
    )
    assert code == 0
    files = sorted(p.name for p in outdir.iterdir())
    assert files == sorted([f"{n}.docx" for n in CERT_NAMES] + ["manifest.csv"])
    for name in CERT_NAMES:
        path = outdir / f"{name}.docx"
        text = "\n".join(p.text for p in Document(path).paragraphs)
        assert name in text
        xml = docx_xml(path)
        assert "{{" not in xml and "}}" not in xml and "{%" not in xml
    assert {r["status"] for r in manifest(outdir)} == {"ok"}


# Criterion 5 -------------------------------------------------------------------------------------------------
def test_placeholders_in_table_header_and_footer(tmp_path):
    data = write_csv(tmp_path / "d.csv", ["Client", "Item", "Amount", "Ref"], [["Asha", "Logo", "1500", "R-9"]])
    template = make_template(
        tmp_path / "t.docx",
        ["Dear {{ client }}"],
        header="Header for {{ client }}",
        footer="Ref {{ ref }}",
        table=[["Item", "Amount"], ["{{ item }}", "{{ amount|money('₹') }}"]],
    )
    outdir = tmp_path / "out"
    assert main(["build", str(data), str(template), "-o", str(outdir)]) == 0
    doc = Document(outdir / "row_1.docx")
    assert doc.tables[0].cell(1, 0).text == "Logo"
    assert doc.tables[0].cell(1, 1).text == "₹1,500.00"
    assert doc.sections[0].header.paragraphs[0].text == "Header for Asha"
    assert doc.sections[0].footer.paragraphs[0].text == "Ref R-9"


def test_conditional_block_and_escaping(tmp_path):
    data = write_csv(tmp_path / "d.csv", ["Name", "VIP"], [["Tom & Jerry <Ltd>", "TRUE"], ["Plain", "FALSE"]])
    template = make_template(tmp_path / "t.docx", ["{{ name }}{% if vip %} (priority){% endif %}"])
    outdir = tmp_path / "out"
    assert main(["build", str(data), str(template), "-o", str(outdir)]) == 0
    assert Document(outdir / "row_1.docx").paragraphs[0].text == "Tom & Jerry <Ltd> (priority)"
    assert Document(outdir / "row_2.docx").paragraphs[0].text == "Plain"


# Criterion 7 -------------------------------------------------------------------------------------------------
def test_duplicate_names_get_suffix_and_check_warns(tmp_path, capsys):
    data = write_csv(tmp_path / "d.csv", ["Name"], [["X"], ["X"], ["Y"]])
    template = make_template(tmp_path / "t.docx", ["{{ name }}"])
    assert main(["check", str(data), str(template), "--name", "{{ name }}"]) == 0
    out = capsys.readouterr().out
    assert any(line.startswith("WARN") and "rows 1, 2" in line and "'X'" in line for line in out.splitlines())
    outdir = tmp_path / "out"
    assert main(["build", str(data), str(template), "-o", str(outdir), "--name", "{{ name }}"]) == 0
    assert {p.name for p in outdir.glob("*.docx")} == {"X.docx", "X_2.docx", "Y.docx"}


def test_name_with_traversal_stays_inside_outdir(tmp_path):
    data = write_csv(tmp_path / "d.csv", ["Name"], [["../evil"], ["../../../tmp/evil2"], ["/abs/evil3"]])
    template = make_template(tmp_path / "t.docx", ["{{ name }}"])
    outdir = tmp_path / "nested" / "out"
    assert main(["build", str(data), str(template), "-o", str(outdir), "--name", "{{ name }}"]) == 0
    written = [p for p in tmp_path.rglob("*.docx") if p.name != "t.docx"]
    assert len(written) == 3
    assert all(p.parent == outdir for p in written)
    assert not (tmp_path / "nested" / "evil.docx").exists()


# Criterion 8 -------------------------------------------------------------------------------------------------
def test_rows_range(examples, tmp_path):
    outdir = tmp_path / "out"
    assert (
        main(
            [
                "build",
                str(examples / "certificates.xlsx"),
                str(examples / "certificate.docx"),
                "-o",
                str(outdir),
                "--rows",
                "2-4",
            ]
        )
        == 0
    )
    assert sorted(p.name for p in outdir.glob("*.docx")) == ["row_2.docx", "row_3.docx", "row_4.docx"]
    statuses = {int(r["row"]): r["status"] for r in manifest(outdir)}
    assert statuses[1] == "skipped" and statuses[3] == "ok" and len(statuses) == 12


def test_where_status_paid(examples, tmp_path):
    outdir = tmp_path / "out"
    assert (
        main(
            [
                "build",
                str(examples / "invoices.csv"),
                str(examples / "invoice.docx"),
                "-o",
                str(outdir),
                "--where",
                "status=paid",
                "--name",
                "{{ invoice_no }}",
            ]
        )
        == 0
    )
    names = sorted(p.stem for p in outdir.glob("*.docx"))
    assert names == ["INV-1001", "INV-1003", "INV-1004", "INV-1006", "INV-1007"]
    for name in names:
        assert "Status: PAID" in "\n".join(p.text for p in Document(outdir / f"{name}.docx").paragraphs)


# Criterion 9 -------------------------------------------------------------------------------------------------
def test_pdf_without_soffice_exits_3_and_keeps_docx(examples, tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("ROWDOCS_SOFFICE", "")  # force "not installed"
    outdir = tmp_path / "out"
    code = main(["build", str(examples / "invoices.csv"), str(examples / "invoice.docx"), "-o", str(outdir), "--pdf"])
    assert code == 3
    err = capsys.readouterr().err
    assert "LibreOffice" in err and "apt install" in err
    assert len(list(outdir.glob("*.docx"))) == 8 and not list(outdir.glob("*.pdf"))
    assert (outdir / "manifest.csv").is_file()


@pytest.mark.skipif(not pdf_works(), reason="LibreOffice (Writer) not available")
def test_pdf_and_merged_pdf(examples, tmp_path):
    outdir = tmp_path / "out"
    code = main(
        [
            "build",
            str(examples / "certificates.xlsx"),
            str(examples / "certificate.docx"),
            "-o",
            str(outdir),
            "--name",
            "{{ certificate_id }}",
            "--pdf",
            "--merge-pdf",
            "--rows",
            "1-4",
        ]
    )
    assert code == 0
    parts = sorted(p for p in outdir.glob("*.pdf") if p.name != "ALL.pdf")
    assert len(parts) == 4
    merged = PdfReader(str(outdir / "ALL.pdf"))
    assert len(merged.pages) == sum(len(PdfReader(str(p)).pages) for p in parts)
    assert "Asha Verma" in merged.pages[0].extract_text()  # row order is kept
    assert all(r["output_pdf"] for r in manifest(outdir) if r["status"] == "ok")


# Criterion 10 ------------------------------------------------------------------------------------------------
def test_zip_contains_every_output_and_manifest(examples, tmp_path):
    outdir = tmp_path / "out"
    assert (
        main(["build", str(examples / "invoices.csv"), str(examples / "invoice.docx"), "-o", str(outdir), "--zip"]) == 0
    )
    with zipfile.ZipFile(tmp_path / "out.zip") as archive:
        assert archive.testzip() is None
        names = set(archive.namelist())
    assert names == {p.name for p in outdir.iterdir()}
    assert "manifest.csv" in names and len([n for n in names if n.endswith(".docx")]) == 8


# Criterion 11 ------------------------------------------------------------------------------------------------
def test_one_bad_row_does_not_stop_the_others(tmp_path, capsys):
    data = write_csv(tmp_path / "d.csv", ["Name", "Amount"], [["A", "10"], ["B", "ten"], ["C", "30"]])
    template = make_template(tmp_path / "t.docx", ["{{ name }}: {{ amount|money('$') }}"])
    outdir = tmp_path / "out"
    assert main(["build", str(data), str(template), "-o", str(outdir)]) == 1
    rows = {int(r["row"]): r for r in manifest(outdir)}
    assert rows[1]["status"] == "ok" and rows[3]["status"] == "ok"
    assert rows[2]["status"] == "error" and "money" in rows[2]["message"]
    assert sorted(p.name for p in outdir.glob("*.docx")) == ["row_1.docx", "row_3.docx"]
    assert "row 2: ERROR" in capsys.readouterr().err


def test_word_smart_quotes_inside_filters_still_work(tmp_path):
    # Word's autocorrect turns money('₹') into money(‘₹’); the template must still render.
    data = write_csv(tmp_path / "d.csv", ["Amount", "Issued"], [["1234567.5", "2026-03-05"]])
    template = make_template(tmp_path / "t.docx", ["{{ amount|money(‘₹’) }} on {{ issued|date(‘%d %b %Y’) }}"])
    outdir = tmp_path / "out"
    assert main(["build", str(data), str(template), "-o", str(outdir)]) == 0
    assert Document(outdir / "row_1.docx").paragraphs[0].text == "₹1,234,567.50 on 05 Mar 2026"
