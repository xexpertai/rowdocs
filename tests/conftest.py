"""Shared fixtures: sample paths, a tiny template builder and a PDF capability probe."""

from __future__ import annotations

import csv
import functools
import io
import zipfile
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
EXAMPLES = ROOT / "examples"


def make_template(
    path: Path,
    body: list[str],
    header: str | None = None,
    footer: str | None = None,
    table: list[list[str]] | None = None,
) -> Path:
    """Write a .docx whose paragraphs/table/header/footer contain the given (Jinja) text."""
    doc = Document()
    for text in body:
        doc.add_paragraph(text)
    if table:
        grid = doc.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    section = doc.sections[0]
    if header is not None:
        section.header.paragraphs[0].text = header
    if footer is not None:
        section.footer.paragraphs[0].text = footer
    doc.save(path)
    return path


def write_csv(path: Path, header: list[str], rows: list[list[str]], delimiter: str = ",") -> Path:
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        writer.writerow(header)
        writer.writerows(rows)
    return path


def docx_xml(path: Path | bytes) -> str:
    """All XML parts of a .docx (body, headers, footers...) joined, for leftover checks."""
    source = io.BytesIO(path) if isinstance(path, bytes) else path
    with zipfile.ZipFile(source) as archive:
        return "\n".join(
            archive.read(n).decode("utf-8") for n in archive.namelist() if n.startswith("word/") and n.endswith(".xml")
        )


@functools.cache
def pdf_works() -> bool:
    """True when LibreOffice is installed *and* can convert a .docx (Writer present)."""
    import tempfile

    from rowdocs.pdf import convert_to_pdf, find_soffice

    if not find_soffice():
        return False
    with tempfile.TemporaryDirectory() as tmp:
        probe = make_template(Path(tmp) / "probe.docx", ["probe"])
        return convert_to_pdf([probe], Path(tmp)).get(probe) is not None


@pytest.fixture
def examples() -> Path:
    return EXAMPLES
