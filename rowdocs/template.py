"""Word template handling: placeholder discovery, rendering and text extraction."""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import parse_xml
from docx.table import Table
from docxtpl import DocxTemplate
from jinja2 import TemplateError, meta

from .filters import make_env

BUILTIN_VARIABLES = frozenset({"_row"})


class TemplateFileError(ValueError):
    """Raised when the template is not a readable .docx file."""


@dataclass
class TemplateInfo:
    variables: dict[str, list[str]] = field(default_factory=dict)  # name -> parts it appears in
    errors: list[str] = field(default_factory=list)  # syntax / unknown-filter problems


class WordTemplate:
    """A .docx template loaded once and rendered many times (once per row)."""

    def __init__(self, source: Path | str | bytes, name: str | None = None):
        if isinstance(source, bytes):
            self.blob = source
            self.name = name or "template.docx"
        else:
            path = Path(source)
            if path.suffix.lower() != ".docx":
                raise TemplateFileError(f"unsupported template {path.name!r}: use a .docx file")
            if not path.is_file():
                raise TemplateFileError(f"template not found: {path}")
            self.blob = path.read_bytes()
            self.name = name or path.name
        if not zipfile.is_zipfile(io.BytesIO(self.blob)):
            raise TemplateFileError(f"{self.name}: not a valid .docx file")
        try:
            Document(io.BytesIO(self.blob))
        except Exception as exc:
            raise TemplateFileError(f"{self.name}: cannot open as .docx ({exc})") from exc

    def _parts(self) -> list[tuple[str, str]]:
        """Return (part name, Jinja source) for the body, headers and footers."""
        tpl = DocxTemplate(io.BytesIO(self.blob))
        doc = Document(io.BytesIO(self.blob))
        parts = [("body", tpl.patch_xml(tpl.xml_to_string(doc.element.body)))]
        for rel in doc.part.rels.values():
            kind = {DocxTemplate.HEADER_URI: "header", DocxTemplate.FOOTER_URI: "footer"}.get(rel.reltype)
            if kind and rel.target_part.blob:
                parts.append((kind, tpl.patch_xml(tpl.xml_to_string(parse_xml(rel.target_part.blob)))))
        return parts

    def inspect(self) -> TemplateInfo:
        """List every placeholder variable (body, tables, headers, footers) and syntax problems."""
        info = TemplateInfo()
        env = make_env()
        for part, source in self._parts():
            try:
                variables = meta.find_undeclared_variables(env.parse(source))
                env.from_string(source)  # compiles: catches unknown filters
            except TemplateError as exc:
                info.errors.append(f"{part}: {exc.__class__.__name__}: {exc.message or exc}")
                continue
            for name in sorted(variables):
                info.variables.setdefault(name, [])
                if part not in info.variables[name]:
                    info.variables[name].append(part)
        return info

    def render(self, context: dict[str, Any]) -> bytes:
        """Render one document and return the .docx bytes."""
        tpl = DocxTemplate(io.BytesIO(self.blob))
        tpl.render(context, jinja_env=make_env(), autoescape=True)
        out = io.BytesIO()
        tpl.save(out)
        return out.getvalue()


def _block_text(block: Any) -> list[str]:
    if isinstance(block, Table):
        lines = []
        for row in block.rows:
            cells = []
            for cell in row.cells:
                text = " ".join(p.text for p in cell.paragraphs).strip()
                if not cells or cells[-1] != text:  # merged cells repeat
                    cells.append(text)
            lines.append(" | ".join(cells))
        return lines
    return [block.text]


def extract_text(docx_bytes: bytes) -> str:
    """Plain text of a document: header, body (paragraphs and tables in order), footer."""
    doc = Document(io.BytesIO(docx_bytes))
    header: list[str] = []
    footer: list[str] = []
    for section in doc.sections:
        for block in section.header.iter_inner_content():
            header.extend(_block_text(block))
        for block in section.footer.iter_inner_content():
            footer.extend(_block_text(block))
    body: list[str] = []
    for block in doc.iter_inner_content():
        body.extend(_block_text(block))
    lines = [t for t in header if t.strip()] + [""] + body + [""] + [t for t in footer if t.strip()]
    return "\n".join(lines).strip("\n")
