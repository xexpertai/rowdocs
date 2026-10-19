"""Load tabular data (.xlsx / .csv), normalise headers and select rows."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

DATA_EXTENSIONS = (".xlsx", ".csv")
_DELIMITERS = (",", ";", "\t")


class DataError(ValueError):
    """Raised when the data file cannot be read or a selection is invalid."""


@dataclass
class Column:
    header: str  # original header text
    key: str  # normalised, template-safe key


@dataclass
class Dataset:
    columns: list[Column]
    rows: list[dict[str, Any]]  # one dict per data row, keyed by Column.key
    source: str = ""
    selected: list[int] = field(default_factory=list)  # 1-based data-row numbers

    @property
    def keys(self) -> list[str]:
        return [c.key for c in self.columns]

    def row(self, number: int) -> dict[str, Any]:
        """Return the row with 1-based number ``number`` (header excluded)."""
        return self.rows[number - 1]

    def context(self, number: int) -> dict[str, Any]:
        """Template context for a row: its values plus the built-in ``_row`` number."""
        return {**self.row(number), "_row": number}


def normalise_header(text: Any) -> str:
    """Turn a column header into a safe template key: "Full Name" -> "full_name"."""
    key = re.sub(r"[^0-9a-z]+", "_", str(text if text is not None else "").strip().lower())
    key = key.strip("_")
    if key and key[0].isdigit():
        key = f"col_{key}"
    return key


def is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _clean_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _csv_value(text: str) -> Any:
    text = text.strip()
    lowered = text.lower()
    if lowered == "true":
        return True
    if lowered == "false":
        return False
    return text


def _build(headers: list[Any], raw_rows: list[list[Any]], source: str) -> Dataset:
    width = max([len(headers)] + [len(r) for r in raw_rows]) if raw_rows or headers else 0
    headers = list(headers) + [None] * (width - len(headers))
    rows = [list(r) + [None] * (width - len(r)) for r in raw_rows]

    keep: list[int] = []
    columns: list[Column] = []
    seen: dict[str, int] = {}
    for index, header in enumerate(headers):
        has_values = any(not is_empty(r[index]) for r in rows)
        if is_empty(header) and not has_values:
            continue  # completely empty column
        key = normalise_header(header) or f"column_{index + 1}"
        if key in seen:
            seen[key] += 1
            key = f"{key}_{seen[key]}"
        else:
            seen[key] = 1
        keep.append(index)
        columns.append(Column(header="" if header is None else str(header).strip(), key=key))

    records = []
    for raw in rows:
        if all(is_empty(v) for v in raw):
            continue  # skip blank lines
        records.append({col.key: _clean_value(raw[i]) for col, i in zip(columns, keep, strict=True)})
    if not columns:
        raise DataError(f"{source}: no header row found")
    return Dataset(columns=columns, rows=records, source=source, selected=list(range(1, len(records) + 1)))


def _detect_delimiter(text: str) -> str:
    first_line = text.split("\n", 1)[0]
    counts = {d: first_line.count(d) for d in _DELIMITERS}
    best = max(_DELIMITERS, key=lambda d: counts[d])  # ties keep the first (",")
    return best if counts[best] else ","


def load_csv_bytes(raw: bytes, source: str = "data.csv") -> Dataset:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DataError(f"{source}: not UTF-8 text. Re-save it as 'CSV UTF-8' and try again.") from exc
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    reader = csv.reader(io.StringIO(text), delimiter=_detect_delimiter(text))
    lines = [line for line in reader]
    if not lines:
        raise DataError(f"{source}: the file is empty")
    return _build(lines[0], [[_csv_value(c) for c in line] for line in lines[1:]], source)


def load_xlsx(path: Path | str, sheet: str | None = None) -> Dataset:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types for corrupt files
        raise DataError(f"{Path(path).name}: cannot open as .xlsx ({exc})") from exc
    try:
        if sheet is None:
            worksheet = workbook.worksheets[0]
        elif sheet in workbook.sheetnames:
            worksheet = workbook[sheet]
        else:
            raise DataError(f"sheet {sheet!r} not found; available: {', '.join(workbook.sheetnames)}")
        values = [list(r) for r in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    if not values:
        raise DataError(f"{Path(path).name}: the sheet is empty")
    return _build(values[0], values[1:], Path(path).name)


def load_data(path: Path | str, sheet: str | None = None) -> Dataset:
    """Load a .xlsx or .csv file into a :class:`Dataset`."""
    path = Path(path)
    if not path.is_file():
        raise DataError(f"data file not found: {path}")
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return load_xlsx(path, sheet)
    if suffix == ".csv":
        return load_csv_bytes(path.read_bytes(), path.name)
    raise DataError(f"unsupported data file {path.name!r}: use .xlsx or .csv")


def parse_rows(spec: str, total: int) -> list[int]:
    """Parse "2-10,15" into sorted 1-based row numbers, checked against ``total``."""
    chosen: set[int] = set()
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        match = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not match:
            raise DataError(f"--rows: cannot read {part!r} (use e.g. 2-10,15)")
        start = int(match.group(1))
        end = int(match.group(2) or start)
        if start < 1 or end < start:
            raise DataError(f"--rows: invalid range {part!r}")
        if end > total:
            raise DataError(f"--rows: row {end} does not exist (the data has {total} rows)")
        chosen.update(range(start, end + 1))
    if not chosen:
        raise DataError("--rows: no rows given")
    return sorted(chosen)


def _as_text(value: Any) -> str:
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime) and value.time() == datetime.min.time():
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def parse_where(expr: str, keys: list[str]) -> tuple[str, bool, str]:
    """Parse "status=paid" / "status!=paid" into (key, negate, value)."""
    match = re.fullmatch(r"\s*([^=!]+?)\s*(!=|=)\s*(.*?)\s*", expr)
    if not match:
        raise DataError(f"--where: cannot read {expr!r} (use column=value or column!=value)")
    key = normalise_header(match.group(1))
    if key not in keys:
        raise DataError(f"--where: unknown column {match.group(1)!r}; columns: {', '.join(keys)}")
    return key, match.group(2) == "!=", match.group(3)


def select_rows(dataset: Dataset, rows: str | None = None, where: list[str] | None = None) -> Dataset:
    """Set ``dataset.selected`` from a --rows spec and --where filters (all must match)."""
    selected = parse_rows(rows, len(dataset.rows)) if rows else list(range(1, len(dataset.rows) + 1))
    for expr in where or []:
        key, negate, wanted = parse_where(expr, dataset.keys)
        wanted_cf = wanted.casefold()
        selected = [
            n for n in selected if (_as_text(dataset.row(n).get(key, "")).strip().casefold() == wanted_cf) != negate
        ]
    dataset.selected = selected
    return dataset
