"""Regenerate the sample files in this folder.

Demo data (fictional): every name, company and amount below is invented.

    python examples/make_examples.py
"""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).parent
NOTE = "Demo data (fictional)"

CERTIFICATES = [
    # certificate id, full name, course, hours, completed on, grade
    ("RSA-2026-001", "Asha Verma", "Spreadsheet Basics", 12, datetime(2026, 3, 5), "Distinction"),
    ("RSA-2026-002", "Tomas Lindqvist", "Spreadsheet Basics", 12, datetime(2026, 3, 5), "Pass"),
    ("RSA-2026-003", "Mei-Ling Ong", "Business Writing", 16, datetime(2026, 3, 12), "Distinction"),
    ("RSA-2026-004", "Kwame Mensah", "Business Writing", 16, datetime(2026, 3, 12), "Pass"),
    ("RSA-2026-005", "Lucia Ferraro", "Data Entry Accuracy", 8, datetime(2026, 3, 19), "Pass"),
    ("RSA-2026-006", "Ravi Kulkarni", "Data Entry Accuracy", 8, datetime(2026, 3, 19), "Distinction"),
    ("RSA-2026-007", "Hana Kobayashi", "Presentation Skills", 10, datetime(2026, 4, 2), "Pass"),
    ("RSA-2026-008", "Daniel O'Connor", "Presentation Skills", 10, datetime(2026, 4, 2), "Pass"),
    ("RSA-2026-009", "Fatima Zahra Idrissi", "Project Planning", 20, datetime(2026, 4, 16), "Distinction"),
    ("RSA-2026-010", "Pedro Alves", "Project Planning", 20, datetime(2026, 4, 16), "Pass"),
    ("RSA-2026-011", "Nadia Petrova", "Customer Support Basics", 6, datetime(2026, 4, 30), "Pass"),
    ("RSA-2026-012", "Samuel Adeyemi", "Customer Support Basics", 6, datetime(2026, 4, 30), "Distinction"),
]

INVOICES = [
    # invoice no, full name, company, issued, item, qty, unit price, amount, status, vip
    ("INV-1001", "Priya Nair", "Lotus Lane Bakery", "2026-03-05", "Menu design", "1", "18500", "18500", "paid", "TRUE"),
    (
        "INV-1002",
        "Arjun Mehta",
        "Copperleaf Tea House",
        "2026-03-07",
        "Logo refresh",
        "1",
        "24000",
        "24000",
        "unpaid",
        "FALSE",
    ),
    (
        "INV-1003",
        "Elena Rossi",
        "Quiet Harbour Books",
        "2026-03-11",
        "Catalogue layout",
        "3",
        "7500",
        "22500",
        "paid",
        "FALSE",
    ),
    (
        "INV-1004",
        "Wei Zhang",
        "Bright Path Tutors",
        "2026-03-14",
        "Course workbook",
        "2",
        "617283.75",
        "1234567.5",
        "paid",
        "TRUE",
    ),
    (
        "INV-1005",
        "Grace Okafor",
        "Green Thread Tailors",
        "2026-03-18",
        "Price list update",
        "1",
        "9000",
        "",
        "unpaid",
        "FALSE",
    ),
    ("INV-1006", "Omar Haddad", "Cedar Step Yoga", "2026-03-21", "Flyer set", "4", "3250", "13000", "paid", "FALSE"),
    (
        "INV-1007",
        "Sofia Martins",
        "Riverstone Ceramics",
        "2026-03-25",
        "Product sheets",
        "6",
        "2100",
        "12600",
        "paid",
        "TRUE",
    ),
    (
        "INV-1008",
        "Ben Carter",
        "Hilltop Repair Co.",
        "2026-03-28",
        "Service brochure",
        "1",
        "15750",
        "15750",
        "unpaid",
        "FALSE",
    ),
]


def _para(
    doc_or_cell,
    text: str,
    size: int = 12,
    bold: bool = False,
    align=None,
    color: RGBColor | None = None,
    space_after: int = 6,
):
    paragraph = doc_or_cell.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    return paragraph


def make_certificates() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Certificates"
    sheet.append(["Certificate ID", "Full Name", "Course", "Hours", "Completed On", "Grade", "Data Note"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in CERTIFICATES:
        sheet.append([*row, NOTE])
    for cell in sheet["E"][1:]:
        cell.number_format = "yyyy-mm-dd"
    for column, width in zip("ABCDEFG", (16, 24, 26, 8, 14, 12, 22), strict=True):
        sheet.column_dimensions[column].width = width
    about = workbook.create_sheet("About")
    about.append([NOTE])
    about.append(["Sample input for rowdocs. Names, courses and dates are invented."])
    workbook.save(HERE / "certificates.xlsx")

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.2))
    navy = RGBColor(0x1F, 0x2A, 0x44)
    grey = RGBColor(0x5B, 0x64, 0x75)
    header = section.header.paragraphs[0]
    header.text = "Riverbend Skills Academy · {{ course }}"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer = section.footer.paragraphs[0]
    footer.text = "Certificate {{ certificate_id }} · {{ data_note }}"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    center = WD_ALIGN_PARAGRAPH.CENTER
    _para(doc, "", space_after=30)
    _para(doc, "Certificate of Completion", 34, True, center, navy, 18)
    _para(doc, "This certifies that", 14, False, center, grey, 8)
    _para(doc, "{{ full_name }}", 30, True, center, navy, 8)
    _para(doc, "has completed the course", 14, False, center, grey, 4)
    _para(doc, "{{ course }} ({{ hours }} hours)", 18, True, center, None, 4)
    _para(doc, "on {{ completed_on|date('%d %B %Y') }}.", 14, False, center, grey, 12)
    _para(
        doc,
        "{% if grade == 'Distinction' %}Awarded with Distinction{% endif %}",
        16,
        True,
        center,
        RGBColor(0x9A, 0x6B, 0x00),
        30,
    )
    _para(doc, "______________________________", 12, False, center, grey, 0)
    _para(doc, "Course Coordinator (fictional)", 11, False, center, grey, 0)
    doc.save(HERE / "certificate.docx")


def make_invoices() -> None:
    with open(HERE / "invoices.csv", "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            [
                "Invoice No",
                "Full Name",
                "Company",
                "Issued",
                "Item",
                "Qty",
                "Unit Price",
                "Amount",
                "Status",
                "VIP",
                "Data Note",
            ]
        )
        for row in INVOICES:
            writer.writerow([*row, NOTE])

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    navy = RGBColor(0x1F, 0x2A, 0x44)
    header = section.header.paragraphs[0]
    header.text = "Invoice {{ invoice_no }} · Studio Paperlane (fictional)"
    footer = section.footer.paragraphs[0]
    footer.text = "{{ data_note }} · Generated with rowdocs"

    _para(doc, "INVOICE", 28, True, None, navy, 2)
    _para(doc, "{{ invoice_no }} · issued {{ issued|date('%d %b %Y') }}", 12, False, None, None, 14)
    _para(doc, "Bill to", 10, True, None, RGBColor(0x5B, 0x64, 0x75), 0)
    _para(doc, "{{ full_name }}", 13, True, None, None, 0)
    _para(doc, "{{ company }}", 12, False, None, None, 14)

    table = doc.add_table(rows=2, cols=4)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ("Item", "Qty", "Unit price", "Amount"), strict=True):
        cell.text = text
    values = ("{{ item }}", "{{ qty }}", "{{ unit_price|money('₹') }}", "{{ amount|money('₹') }}")
    for cell, text in zip(table.rows[1].cells, values, strict=True):
        cell.text = text

    _para(doc, "", space_after=6)
    _para(doc, "Total due: {{ amount|money('₹') }}", 14, True, None, navy, 4)
    _para(doc, "Status: {{ status|upper }}", 12, False, None, None, 10)
    _para(
        doc,
        "{% if vip %}Priority client: thank you for working with us again this year.{% endif %}",
        11,
        False,
        None,
        RGBColor(0x1D, 0x7A, 0x46),
        10,
    )
    _para(doc, "Payment terms: 15 days. Please quote the invoice number with your payment.", 10)
    doc.save(HERE / "invoice.docx")


if __name__ == "__main__":
    make_certificates()
    make_invoices()
    print(f"wrote samples to {HERE}")
